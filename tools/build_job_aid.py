#!/usr/bin/env python3
"""Generate docs/Whitespace_Analyzer_Job_Aid.pptx — 1-page weekly update job aid."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── SAP Classic Corporate colors ─────────────────────────────────────────────────
DARK_NAVY  = RGBColor(0x00, 0x2A, 0x86)
DEEP_NAVY  = RGBColor(0x00, 0x14, 0x4A)
SAP_BLUE   = RGBColor(0x1B, 0x90, 0xFF)
LIGHT_BLUE = RGBColor(0x89, 0xD1, 0xFF)
LIGHT_SKY  = RGBColor(0xD1, 0xEF, 0xFF)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY  = RGBColor(0x44, 0x44, 0x44)
MED_GRAY   = RGBColor(0x88, 0x88, 0x88)
LIGHT_GRAY = RGBColor(0xCC, 0xD4, 0xE0)
PALE_BG    = RGBColor(0xF5, 0xF8, 0xFF)
CODE_BG    = RGBColor(0xE8, 0xEC, 0xF5)
GREEN_OK   = RGBColor(0x10, 0x7C, 0x10)
AMBER_WARN = RGBColor(0x8B, 0x50, 0x00)
RED_FAIL   = RGBColor(0x9B, 0x10, 0x10)
GREEN_BG   = RGBColor(0xE8, 0xF5, 0xE9)
AMBER_BG   = RGBColor(0xFF, 0xF3, 0xCD)
RED_BG     = RGBColor(0xFD, 0xED, 0xED)

F  = "Arial"
FC = "Courier New"

# ── Low-level helpers ─────────────────────────────────────────────────────────────

def add_rect(slide, l, t, w, h, fill=None, line=None, lw=0.5):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s

def add_oval(slide, l, t, w, h, fill=None):
    s = slide.shapes.add_shape(9, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    s.line.fill.background()
    tf = s.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    return s

def add_tb(slide, l, t, w, h, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = wrap
    return tb.text_frame

def rn(p, text, sz=9, bold=False, color=BLACK, font=F, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r

def fp(tf, text, sz=9, bold=False, color=BLACK, font=F, align=PP_ALIGN.LEFT, sb=0, sa=0, italic=False):
    p = tf.paragraphs[0]
    p.alignment = align
    if sb: p.space_before = Pt(sb)
    if sa: p.space_after = Pt(sa)
    if text:
        rn(p, text, sz=sz, bold=bold, color=color, font=font, italic=italic)
    return p

def ap(tf, text, sz=9, bold=False, color=BLACK, font=F, align=PP_ALIGN.LEFT, sb=0, sa=0, italic=False):
    p = tf.add_paragraph()
    p.alignment = align
    if sb: p.space_before = Pt(sb)
    if sa: p.space_after = Pt(sa)
    if text:
        rn(p, text, sz=sz, bold=bold, color=color, font=font, italic=italic)
    return p

# ── Component helpers ─────────────────────────────────────────────────────────────

def col_header(slide, l, t, w, label):
    add_rect(slide, l, t, w, 0.30, fill=DARK_NAVY)
    tf = add_tb(slide, l+0.14, t+0.05, w-0.25, 0.22, wrap=False)
    fp(tf, label, sz=7.5, bold=True, color=WHITE)

def step_label(slide, l, t, num, label):
    ov = add_oval(slide, l, t+0.01, 0.22, 0.22, fill=SAP_BLUE)
    p = ov.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    rn(p, str(num), sz=8, bold=True, color=WHITE)
    tf = add_tb(slide, l+0.28, t, 3.5, 0.22, wrap=False)
    fp(tf, label, sz=9.5, bold=True, color=DARK_NAVY)

def check_item(slide, l, t, w, label, filename=None, filter_text=None):
    """Checkbox square + bold label + optional filename + filter line."""
    cb = slide.shapes.add_shape(1, Inches(l), Inches(t+0.02), Inches(0.13), Inches(0.13))
    cb.fill.background()
    cb.line.color.rgb = SAP_BLUE
    cb.line.width = Pt(1.0)

    tf = add_tb(slide, l+0.19, t, w-0.19, 0.19, wrap=False)
    fp(tf, label, sz=8.5, bold=True, color=BLACK)

    row = 0.21
    if filename:
        tf2 = add_tb(slide, l+0.19, t+row, w-0.19, 0.16, wrap=False)
        fp(tf2, filename, sz=7, color=DARK_GRAY, italic=True, font=FC)
        row += 0.17

    if filter_text:
        tf3 = add_tb(slide, l+0.19, t+row, w-0.19, 0.16, wrap=True)
        p = tf3.paragraphs[0]
        rn(p, "Filter: ", sz=7, bold=True, color=MED_GRAY)
        rn(p, filter_text, sz=7, color=MED_GRAY)

def code_box(slide, l, t, w, h, lines, fsz=7.5):
    add_rect(slide, l, t, w, h, fill=CODE_BG, line=LIGHT_GRAY, lw=0.5)
    tf = add_tb(slide, l+0.1, t+0.055, w-0.18, h-0.1)
    fp(tf, lines[0], sz=fsz, font=FC, color=DEEP_NAVY)
    for ln in lines[1:]:
        ap(tf, ln, sz=fsz, font=FC, color=DEEP_NAVY)

def note_box(slide, l, t, w, h, label, text, fill=LIGHT_SKY):
    add_rect(slide, l, t, w, h, fill=fill, line=None)
    tf = add_tb(slide, l+0.1, t+0.06, w-0.18, h-0.1)
    p = tf.paragraphs[0]
    rn(p, label, sz=7.5, bold=True, color=DARK_NAVY)
    rn(p, text, sz=7.5, color=DARK_GRAY)

def status_badge(slide, l, t, w, h, prefix, code_lbl, desc, fill, fg):
    add_rect(slide, l, t, w, h, fill=fill, line=LIGHT_GRAY, lw=0.5)
    tf = add_tb(slide, l+0.1, t+0.05, w-0.18, h-0.1)
    p = tf.paragraphs[0]
    rn(p, prefix + "  ", sz=9, bold=True, color=fg)
    rn(p, code_lbl, sz=8.5, bold=True, color=fg)
    rn(p, "  " + desc, sz=8, color=DARK_GRAY)

def divider(slide, l, t, w):
    add_rect(slide, l, t, w, 0.01, fill=LIGHT_GRAY)

# ── Build presentation ─────────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# ── HEADER ────────────────────────────────────────────────────────────────────────
add_rect(slide, 0, 0, 13.33, 0.9, fill=DARK_NAVY)

tf = add_tb(slide, 0.35, 0.08, 12.5, 0.42, wrap=False)
p = tf.paragraphs[0]
rn(p, "Whitespace Analyzer", sz=20, color=WHITE)
rn(p, "  ·  Weekly Update Job Aid", sz=20, color=LIGHT_BLUE)

tf2 = add_tb(slide, 0.35, 0.52, 10.5, 0.28)
fp(tf2, "Run each week to refresh the dashboard with the latest account coverage, pipeline, and outside-in opportunity data.",
   sz=8, color=LIGHT_SKY)

# ── COLUMNS ───────────────────────────────────────────────────────────────────────
CT = 0.96   # column top
CH = 6.2    # column height
CW = 4.1    # column width
C1, C2, C3 = 0.22, 4.57, 8.92
P = 0.18    # inner padding

for cx in (C1, C2, C3):
    add_rect(slide, cx, CT, CW, CH, fill=PALE_BG, line=LIGHT_GRAY, lw=0.5)

col_header(slide, C1, CT, CW, "STEPS 1–2   ·   SETUP & DOWNLOADS")
col_header(slide, C2, CT, CW, "STEP 3   ·   EXPORT TOP PARENT IDs")
col_header(slide, C3, CT, CW, "STEPS 4–5   ·   RUN & VALIDATE")

# ── COLUMN 1: Setup & Downloads ───────────────────────────────────────────────────
y = CT + 0.38

step_label(slide, C1+P, y, 1, "Create the input folder")
y += 0.28

tf = add_tb(slide, C1+P, y, CW-P*2, 0.28)
fp(tf, "Name the folder with today’s data-pull date:", sz=8, color=DARK_GRAY)
y += 0.28

code_box(slide, C1+P, y, CW-P*2, 0.25, ["input/2026-08-13/"])
y += 0.32

divider(slide, C1+P, y, CW-P*2)
y += 0.12

step_label(slide, C1+P, y, 2, "Download these reports into that folder")
y += 0.3

# Checklist items
check_item(slide, C1+P, y, CW-P*2,
    "Adoption Monitor (AM)",
    "MW_MU_AM_YYYY-MM-DD.csv",
    "SBB Market Unit Desc = US Midwest")
y += 0.56

check_item(slide, C1+P, y, CW-P*2,
    "Account Master (IAA)",
    "MW_MU_IAA_YYYY-MM-DD.csv",
    "Market Unit = Midwest")
y += 0.56

check_item(slide, C1+P, y, CW-P*2,
    "Pipeline Forecast (ADRM)",
    "MW_MU_ADRM_YYYY-MM-DD.csv",
    "Zero Flag=N · Source=CLOUD · Region L1=Americas")
y += 0.56

check_item(slide, C1+P, y, CW-P*2,
    "Outside-In Wallet (TAO)  —  3 files",
    "US_MIDWEST_TAO_DD_MM_YYYY_v2_1/2/3.csv")
y += 0.42

check_item(slide, C1+P, y, CW-P*2,
    "Planned Pipeline (SPI)  —  optional",
    "MW_MU_SPI_YYYY-MM-DD.csv",
    "Region L2=US Midwest · Currency=USD")

# ── COLUMN 2: Export Top Parent IDs ──────────────────────────────────────────────
y = CT + 0.38

step_label(slide, C2+P, y, 3, "Generate the Top Parent ID list")
y += 0.3

tf = add_tb(slide, C2+P, y, CW-P*2, 0.3)
fp(tf, "After downloading IAA (before downloading AM-TP), run:", sz=8, color=DARK_GRAY)
y += 0.3

code_box(slide, C2+P, y, CW-P*2, 0.42, [
    "python3 tools/export_top_parent_ids.py \\",
    "  --iaa input/DATE/MW_MU_IAA_DATE.csv",
])
y += 0.5

tf = add_tb(slide, C2+P, y, CW-P*2, 0.46)
p = tf.paragraphs[0]
rn(p, "Writes ", sz=8, color=DARK_GRAY)
rn(p, "top_parent_ids.csv", sz=8, color=DEEP_NAVY, font=FC)
rn(p, " into the dated input folder. Use this list to scope the AM-TP export in Adoption Monitor.", sz=8, color=DARK_GRAY)
y += 0.5

divider(slide, C2+P, y, CW-P*2)
y += 0.12

tf3 = add_tb(slide, C2+P, y, CW-P*2, 0.2)
fp(tf3, "Then download and add to the same folder:", sz=8, bold=True, color=DARK_NAVY)
y += 0.25

check_item(slide, C2+P, y, CW-P*2,
    "AM Top Parent (AM-TP)",
    "MW_MU_AM_TP_YYYY-MM-DD.csv",
    "Planning Entity IDs in top_parent_ids.csv")
y += 0.58

note_box(slide, C2+P, y, CW-P*2, 0.38,
    "Note: ", "AM-TP is optional — the pipeline runs without it. Skip if unavailable.")
y += 0.46

divider(slide, C2+P, y, CW-P*2)
y += 0.12

tf4 = add_tb(slide, C2+P, y, CW-P*2, 0.2)
fp(tf4, "All files go in one dated folder:", sz=8, bold=True, color=DARK_NAVY)
y += 0.24

code_box(slide, C2+P, y, CW-P*2, 0.88, [
    "input/",
    "  2026-08-13/",
    "    MW_MU_AM_2026-08-13.csv",
    "    MW_MU_IAA_2026-08-13.csv",
    "    MW_MU_ADRM_2026-08-13.csv",
    "    US_MIDWEST_TAO_13_08_2026_v2_1.csv",
    "    US_MIDWEST_TAO_13_08_2026_v2_2.csv",
    "    US_MIDWEST_TAO_13_08_2026_v2_3.csv",
    "    MW_MU_AM_TP_2026-08-13.csv  ← optional",
], fsz=7.0)

# ── COLUMN 3: Run & Validate ──────────────────────────────────────────────────────
y = CT + 0.38

step_label(slide, C3+P, y, 4, "Run the pipeline")
y += 0.3

tf = add_tb(slide, C3+P, y, CW-P*2, 0.28)
fp(tf, "Point the pipeline at the dated input folder:", sz=8, color=DARK_GRAY)
y += 0.28

code_box(slide, C3+P, y, CW-P*2, 0.42, [
    "python3 run_pipeline.py \\",
    "  --input-dir input/2026-08-13/",
])
y += 0.5

tf = add_tb(slide, C3+P, y, CW-P*2, 0.28)
fp(tf, "The folder name sets the snapshot date shown in the dashboard header.", sz=8, color=DARK_GRAY)
y += 0.35

divider(slide, C3+P, y, CW-P*2)
y += 0.12

step_label(slide, C3+P, y, 5, "Check validation results")
y += 0.32

status_badge(slide, C3+P, y, CW-P*2, 0.30,
    "✓", "Exit 0", "Clean run — safe to share", GREEN_BG, GREEN_OK)
y += 0.35

status_badge(slide, C3+P, y, CW-P*2, 0.30,
    "⚠", "Exit 2", "Warnings — review printed report before sharing", AMBER_BG, AMBER_WARN)
y += 0.35

status_badge(slide, C3+P, y, CW-P*2, 0.30,
    "✗", "Exit 1", "Failed — fix the error and re-run", RED_BG, RED_FAIL)
y += 0.4

divider(slide, C3+P, y, CW-P*2)
y += 0.12

tf_out = add_tb(slide, C3+P, y, CW-P*2, 0.2)
fp(tf_out, "Output — open in any browser:", sz=8, bold=True, color=DARK_NAVY)
y += 0.25

code_box(slide, C3+P, y, CW-P*2, 0.42, [
    "output/2026-08-13/",
    "  whitespace-analyzer.html   ← open this",
])
y += 0.5

note_box(slide, C3+P, y, CW-P*2, 0.45,
    "Common exit 2 causes: ",
    "AM↔IAA match rate below threshold · TAO coverage below 50% · Missing expected column")

# ── FOOTER ────────────────────────────────────────────────────────────────────────
tf_f = add_tb(slide, 0.35, 7.3, 4, 0.15, wrap=False)
fp(tf_f, "SAP Internal", sz=8, color=MED_GRAY)

# ── SAVE ──────────────────────────────────────────────────────────────────────────
out = "docs/Whitespace_Analyzer_Job_Aid.pptx"
prs.save(out)
print(f"Saved: {out}")
